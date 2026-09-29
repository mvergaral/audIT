#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Planilla de cálculo del mandante y partidas de la oferta económica.

  python3 herramientas/planilla.py comparar [--muestra]
        compara cada partida de economico/datos/partidas.csv que declara su
        celda (columna planilla, como Resumen!C5) con el valor de esa celda
        en la planilla de economico/planilla/
  python3 herramientas/planilla.py volcar [--muestra] SALIDA.xlsx
        escribe en una copia de la planilla el valor neto en pesos de cada
        partida en su celda. No toca las celdas con fórmula: avisa.

La planilla la entrega el mandante (Formulario T-22, FEP01 p.69) y es el
modelo financiero del Sobre N.º 3 (Formulario E-21, entregable 3). El formato
no la genera. El Formulario E-21 exige que los valores de los tres documentos
sean idénticos y declara la discrepancia causal de descalificación: por eso
las cifras salen de una sola fuente, partidas.csv, que alimenta los PDF, el
Markdown, el DOCX y, con esta herramienta, la planilla.

Usa openpyxl del entorno de herramientas/.venv (herramientas/preparar.sh).
"""
import glob
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
from comun import RAIZ, Config  # noqa: E402

try:
    import openpyxl
except ImportError:              # compilar.py importa este módulo sin el entorno
    openpyxl = None


def al_entorno():
    """Vuelve a lanzar esta herramienta con el Python de herramientas/.venv."""
    venv = os.path.join(AQUI, ".venv")
    if os.path.exists(os.path.join(venv, "bin", "python")) and \
            os.path.realpath(sys.prefix) != os.path.realpath(venv):
        py = os.path.join(venv, "bin", "python")
        os.execv(py, [py] + sys.argv)

RAICES = ["economico"]
RAICES_MUESTRA = ["muestra/economico", "economico"]


# ----------------------------------------------------------------------------
#  Partidas (misma lectura y misma aritmética entera que estilo/economico.lua)
# ----------------------------------------------------------------------------
def div_red(a, b):
    if b < 0:
        a, b = -a, -b
    if a >= 0:
        return (2 * a + b) // (2 * b)
    return -((2 * (-a) + b) // (2 * b))


def leer_numero(s):
    s = (s or "").replace(" ", "")
    if not s:
        return None
    neg = s.startswith("-")
    s = s.lstrip("-")
    ent, _, dec = s.replace(",", ".").partition(".")
    if not ent.isdigit() or (dec and not dec.isdigit()):
        raise ValueError(f"«{s}» no es un número")
    num, den = int(ent + dec), 10 ** len(dec)
    return (-num if neg else num), den


def leer_partidas(raices):
    """clave -> dict. Si una clave se repite entre raíces, gana la primera."""
    partidas = {}
    for r in raices:
        ruta = os.path.join(RAIZ, r, "datos", "partidas.csv")
        if not os.path.exists(ruta):
            continue
        for n, linea in enumerate(open(ruta, encoding="utf-8")):
            linea = linea.rstrip("\r\n")
            if n == 0 or not linea.strip() or linea.lstrip().startswith("#"):
                continue
            c = [x.strip() for x in linea.split(";")] + [""] * 6
            if c[0] in partidas:
                continue
            partidas[c[0]] = {"clave": c[0], "concepto": c[1], "moneda": c[2] or "CLP",
                              "valor": c[3], "iva": c[4].lower() != "no", "planilla": c[5]}
    return partidas


def neto_clp(partidas, clave, tc, pila=()):
    p = partidas[clave]
    if p["valor"].startswith("="):
        if clave in pila:
            raise ValueError(f"la suma de {clave} se incluye a sí misma")
        total = 0
        for k in p["valor"][1:].replace(" ", "").split("+"):
            v = neto_clp(partidas, k, tc, pila + (clave,))
            if v is None:
                return None
            total += v
        return total
    n = leer_numero(p["valor"])
    if n is None:
        return None
    num, den = n
    tasa = 1 if p["moneda"] == "CLP" else tc[p["moneda"]]
    return div_red(num * tasa, den)


# ----------------------------------------------------------------------------
#  Planilla
# ----------------------------------------------------------------------------
def buscar_planilla(raices):
    """El único .xlsx de <raíz>/planilla/, o None."""
    for r in raices:
        xs = sorted(glob.glob(os.path.join(RAIZ, r, "planilla", "*.xlsx")))
        xs = [x for x in xs if not os.path.basename(x).startswith("~$")]
        if len(xs) > 1:
            raise SystemExit(f"Hay más de una planilla en {r}/planilla/: {', '.join(map(os.path.basename, xs))}")
        if xs:
            return xs[0]
    return None


def celda(libro, ref):
    hoja, _, dir_ = ref.rpartition("!")
    hoja = hoja.strip("'")
    if hoja not in libro.sheetnames:
        raise KeyError(f"la planilla no tiene la hoja «{hoja}»")
    return libro[hoja][dir_]


def comparar(xlsx, partidas, tc):
    """Lista de (clave, celda, valor en partidas, valor en la planilla, ok)."""
    if openpyxl is None:
        raise SystemExit("Falta openpyxl: correr herramientas/preparar.sh")
    libro = openpyxl.load_workbook(xlsx, data_only=True)
    salida = []
    for clave, p in partidas.items():
        if not p["planilla"]:
            continue
        esperado = neto_clp(partidas, clave, tc)
        try:
            v = celda(libro, p["planilla"]).value
        except KeyError as e:
            salida.append((clave, p["planilla"], esperado, str(e), False))
            continue
        if v is None and esperado is not None and _es_formula(xlsx, p["planilla"]):
            v = "fórmula sin valor calculado: abrir y guardar la planilla en Excel"
        if esperado is None:
            ok = v in (None, "")
        else:
            try:
                ok = v is not None and round(float(v)) == esperado
            except (TypeError, ValueError):
                ok = False
        salida.append((clave, p["planilla"], esperado, v, ok))
    return salida


def _es_formula(xlsx, ref):
    c = celda(openpyxl.load_workbook(xlsx), ref)
    return isinstance(c.value, str) and c.value.startswith("=")


def volcar(xlsx, destino, partidas, tc):
    if openpyxl is None:
        raise SystemExit("Falta openpyxl: correr herramientas/preparar.sh")
    libro = openpyxl.load_workbook(xlsx)
    avisos = []
    for clave, p in partidas.items():
        if not p["planilla"]:
            continue
        v = neto_clp(partidas, clave, tc)
        c = celda(libro, p["planilla"])
        if isinstance(c.value, str) and c.value.startswith("="):
            avisos.append(f"{p['planilla']} tiene una fórmula: no se escribe {clave}")
            continue
        if v is not None:
            c.value = v
    libro.save(destino)
    return avisos


def main():
    if openpyxl is None:
        al_entorno()
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args or args[0] not in ("comparar", "volcar"):
        print(__doc__)
        return 2
    raices = RAICES_MUESTRA if "--muestra" in sys.argv else RAICES
    cfg = Config()
    partidas = leer_partidas(raices)
    xlsx = buscar_planilla(raices)
    if not xlsx:
        print("No hay planilla en " + " ni en ".join(f"{r}/planilla/" for r in raices)
              + ". La entrega el mandante.")
        return 1
    if args[0] == "comparar":
        malos = 0
        for clave, ref, esperado, v, ok in comparar(xlsx, partidas, cfg.tipos_cambio):
            malos += not ok
            print(f"{'igual   ' if ok else 'DISTINTO'} {clave:32s} {ref:18s} partidas {esperado}  planilla {v}")
        sin_celda = [k for k, p in partidas.items() if not p["planilla"]]
        if sin_celda:
            print(f"{len(sin_celda)} partidas sin celda asignada en la columna planilla")
        return 1 if malos else 0
    destino = args[1] if len(args) > 1 else None
    if not destino:
        print("Falta el archivo de salida")
        return 2
    for a in volcar(xlsx, destino, partidas, cfg.tipos_cambio):
        print("aviso: " + a)
    print("escrito " + destino)
    return 0


if __name__ == "__main__":
    sys.exit(main())
